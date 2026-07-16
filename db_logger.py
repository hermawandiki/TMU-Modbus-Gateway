import threading
import time
import os
import logging
import openpyxl
from openpyxl import Workbook
from influxdb import InfluxDBClient
from data_handler import lcd_lock, lcd_values, lcd_fields

# Influx Config
INFLUX_HOST = 'localhost'; INFLUX_PORT = 8086; INFLUX_DB = 'all_sensor_data'
INFLUX_USER = ''; INFLUX_PASS = ''

FIELDS = [
    "tsu_oil_level", "tsu_oil_temp", "tsu_oil_pressure",
    "edmcr_oil_level", "edmcr_oil_temp", "edmcr_oil_pressure",
    "bus_u_temp", "bus_v_temp", "bus_w_temp",
    "wit_u_temp", "wit_v_temp", "wit_w_temp",
    "u_n_volt", "v_n_volt", "w_n_volt",
    "u_v_volt", "v_w_volt", "u_w_volt",
    "u_current", "v_current", "w_current", "avg_current"
]

def _excel_loop(stop_event):
    log_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logsys")
    os.makedirs(log_dir, exist_ok=True)
    while not stop_event.is_set():
        try:
            now = time.localtime()
            file_path = os.path.join(log_dir, f"datalog_{time.strftime('%Y_%m', now)}.xlsx")
            sheet_name = time.strftime("%Y-%m-%d", now)
            timestamp = time.strftime("%H:%M:%S", now)
            with lcd_lock: val_copy = lcd_values.copy()
            
            wb = openpyxl.load_workbook(file_path) if os.path.exists(file_path) else Workbook()
            if "Sheet" in wb.sheetnames: wb.remove(wb["Sheet"])
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(title=sheet_name)
            if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
                ws.append(["Jam"] + lcd_fields)
            ws.append([timestamp] + val_copy)
            wb.save(file_path)
        except Exception as e: logging.error(f"Excel Error: {e}")
        for _ in range(10):
            if stop_event.is_set(): break
            time.sleep(1)

def _influx_loop(stop_event, hb_callback):
    client = InfluxDBClient(host=INFLUX_HOST, port=INFLUX_PORT, username=INFLUX_USER, password=INFLUX_PASS, database=INFLUX_DB)
    while not stop_event.is_set():
        try:
            with lcd_lock: val_copy = lcd_values.copy()
            if len(val_copy) == len(FIELDS):
                json_body = [{"measurement": "tmu_monitoring", "tags": {"location": "Surabaya", "device_id": "TMU-01"}, "fields": {}}]
                for i, fn in enumerate(FIELDS): json_body[0]["fields"][fn] = float(val_copy[i])
                client.write_points(json_body)
                if hb_callback: hb_callback(time.time())
        except Exception: pass
        for _ in range(2):
            if stop_event.is_set(): break
            time.sleep(1)
    try: client.close()
    except: pass

def start_db_logger_engine(stop_event, hb_callback=None):
    threading.Thread(target=_excel_loop, args=(stop_event,), daemon=True).start()
    threading.Thread(target=_influx_loop, args=(stop_event, hb_callback), daemon=True).start()